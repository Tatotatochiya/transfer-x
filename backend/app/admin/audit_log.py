"""The Audit log page and its Excel export (admin panel).

Reads every audit event, not only admin ones: deal, offer, loan and
preference events sit beside staff actions, so a dispute can be traced in
one place. `admin_only` narrows it to what TransferX staff did.
"""
import io
import uuid
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone

from sqlalchemy import String, cast, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.admin.audit import ADMIN_PREFIX
from app.audit.models import AuditEvent
from app.auth.models import User

EXPORT_LIMIT = 50_000

# Where an entity can be opened from the log.
ENTITY_LINKS = {
    "club": "/admin/clubs/{id}",
    "player": "/admin/players/{id}",
    "deal": "/deals/{id}",
    "DEAL": "/deals/{id}",
    "offer": "/offers/{id}",
    "OFFER": "/offers/{id}",
    "sale": "/sales/{id}",
}


@dataclass
class AuditFilters:
    q: str | None = None
    action: str | None = None
    entity_type: str | None = None
    actor_id: uuid.UUID | None = None
    entity_id: uuid.UUID | None = None
    admin_only: bool = False
    date_from: date | None = None
    date_to: date | None = None

    def describe(self) -> list[tuple[str, str]]:
        rows = []
        for label, value in (
            ("Search", self.q), ("Action", self.action), ("Entity type", self.entity_type),
            ("Actor", self.actor_id), ("Entity", self.entity_id),
            ("From", self.date_from), ("To", self.date_to),
        ):
            if value:
                rows.append((label, str(value)))
        if self.admin_only:
            rows.append(("Only", "TransferX staff actions"))
        return rows or [("Filters", "None: every event")]


def _query(f: AuditFilters):
    q = select(AuditEvent, User.email, User.first_name, User.last_name).outerjoin(User, User.id == AuditEvent.actor_user_id)
    if f.admin_only:
        q = q.where(AuditEvent.action.startswith(ADMIN_PREFIX))
    if f.action:
        q = q.where(AuditEvent.action == f.action)
    if f.entity_type:
        q = q.where(AuditEvent.entity_type == f.entity_type)
    if f.actor_id:
        q = q.where(AuditEvent.actor_user_id == f.actor_id)
    if f.entity_id:
        q = q.where(AuditEvent.entity_id == f.entity_id)
    if f.date_from:
        q = q.where(AuditEvent.created_at >= datetime.combine(f.date_from, time.min, tzinfo=timezone.utc))
    if f.date_to:
        q = q.where(AuditEvent.created_at < datetime.combine(f.date_to + timedelta(days=1), time.min, tzinfo=timezone.utc))
    if f.q:
        like = f"%{f.q.strip()}%"
        q = q.where(or_(
            AuditEvent.description.ilike(like), AuditEvent.action.ilike(like), User.email.ilike(like),
            User.first_name.ilike(like), User.last_name.ilike(like),
            cast(AuditEvent.entity_id, String).ilike(like),
        ))
    return q


def _name(first: str | None, last: str | None) -> str | None:
    return " ".join(p for p in (first, last) if p) or None


def _row(event: AuditEvent, email: str | None, first: str | None = None, last: str | None = None) -> dict:
    payload = event.payload_json or {}
    link = ENTITY_LINKS.get(event.entity_type)
    return {
        "id": event.id,
        "created_at": event.created_at,
        "actor_user_id": event.actor_user_id,
        "actor_email": email,
        "actor_name": _name(first, last),
        "action": event.action,
        "entity_type": event.entity_type,
        "entity_id": event.entity_id,
        "description": event.description,
        "reason": payload.get("reason"),
        "payload": payload,
        "link": link.format(id=event.entity_id) if link else None,
        "by_staff": event.action.startswith(ADMIN_PREFIX),
    }


async def list_events(db: AsyncSession, f: AuditFilters, page: int, page_size: int) -> tuple[list[dict], int]:
    q = _query(f)
    total = (await db.execute(select(func.count()).select_from(q.subquery()))).scalar_one()
    rows = (await db.execute(
        q.order_by(AuditEvent.created_at.desc()).offset((page - 1) * page_size).limit(page_size)
    )).all()
    return [_row(e, email, first, last) for e, email, first, last in rows], total


async def facets(db: AsyncSession) -> dict:
    actions = (await db.execute(select(AuditEvent.action).distinct().order_by(AuditEvent.action))).scalars().all()
    types = (await db.execute(select(AuditEvent.entity_type).distinct().order_by(AuditEvent.entity_type))).scalars().all()
    return {"actions": list(actions), "entity_types": list(types)}


def _readable(action: str) -> str:
    """'admin.user.updated' → 'User updated (staff)'."""
    staff = action.startswith(ADMIN_PREFIX)
    words = action.removeprefix(ADMIN_PREFIX).replace(".", " ").replace("_", " ").strip()
    words = words[:1].upper() + words[1:].lower()
    return f"{words} (staff)" if staff else words


async def export_xlsx(db: AsyncSession, f: AuditFilters, *, exported_by: str) -> tuple[bytes, int, bool]:
    """The filtered log as an Excel workbook: an "Audit log" sheet (newest
    first, filterable header, frozen top row) and an "About" sheet with who
    exported it, when, and the filters. Returns (bytes, rows, truncated)."""
    import json

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = (await db.execute(
        _query(f).order_by(AuditEvent.created_at.desc()).limit(EXPORT_LIMIT + 1)
    )).all()
    truncated = len(rows) > EXPORT_LIMIT
    rows = rows[:EXPORT_LIMIT]

    wb = Workbook()
    ws = wb.active
    ws.title = "Audit log"
    headers = ["When (UTC)", "Who", "What", "Action code", "Entity type", "Entity ID", "Description", "Reason", "Details"]
    widths = [20, 32, 34, 30, 16, 38, 60, 40, 60]
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    head_fill = PatternFill("solid", fgColor="EAF1FB")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = head_fill
    for event, email, first, last in rows:
        name = _name(first, last)
        payload = dict(event.payload_json or {})
        reason = payload.pop("reason", None)
        payload.pop("admin_action", None)
        when = event.created_at
        if when.tzinfo is not None:
            when = when.astimezone(timezone.utc).replace(tzinfo=None)
        ws.append([
            when, (f"{name} ({email})" if name and email else email)
            or ("System" if event.actor_user_id is None else str(event.actor_user_id)),
            _readable(event.action), event.action, event.entity_type, str(event.entity_id),
            event.description or "", reason or "",
            json.dumps(payload, ensure_ascii=False, default=str) if payload else "",
        ])
    for cell in ws["A"][1:]:
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    for col in ("G", "H", "I"):
        for cell in ws[col][1:]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"

    about = wb.create_sheet("About")
    about.column_dimensions["A"].width = 22
    about.column_dimensions["B"].width = 70
    about.append(["TransferX audit log export"])
    about["A1"].font = Font(bold=True, size=13)
    about.append([])
    for label, value in [
        ("Exported", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Exported by", exported_by),
        ("Rows", f"{len(rows):,}" + (f" (first {EXPORT_LIMIT:,} only; narrow the filters for the rest)" if truncated else "")),
        *f.describe(),
    ]:
        about.append([label, value])
        about.cell(row=about.max_row, column=1).font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows), truncated
