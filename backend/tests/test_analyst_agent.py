"""The analyst loop (AI analyst spec §3, §5, §7), with a scripted model."""
import json
from types import SimpleNamespace

import pytest
from httpx import AsyncClient

from app.config import settings
from tests.conftest import _auth_headers
from tests.test_analyst_tools import _setup

pytestmark = pytest.mark.asyncio


def _call(name, args, i=0):
    return SimpleNamespace(id=f"c{i}", function=SimpleNamespace(name=name, arguments=json.dumps(args)))


@pytest.fixture
def model(monkeypatch):
    """The model's turns, in order; records what it was sent."""
    from app.ai import client

    monkeypatch.setattr(settings, "anthropic_api_key", "test-key")
    script: list = []
    seen: list = []

    async def fake(messages, tools, **kw):
        seen.append({"messages": [dict(m) for m in messages], "tools": [t["function"]["name"] for t in tools], **kw})
        return script.pop(0) if script else SimpleNamespace(content="Done.", tool_calls=None)

    monkeypatch.setattr(client, "chat_with_tools", fake)
    return script, seen


def _turn(*calls):
    return SimpleNamespace(content="", tool_calls=list(calls))


async def _ask(client, club, question, history=None):
    resp = await client.post("/ai/analyst", json={"question": question, "history": history or []}, headers=_auth_headers(club))
    assert resp.status_code == 200, resp.text
    return resp.json()


async def test_five_listed_midfielders_comes_back_as_a_table_built_from_the_tool(client: AsyncClient, db, model):
    script, seen = model
    seller, buyer, mid, *_ = await _setup(client, db)
    script += [
        _turn(_call("search_players", {"position": "MID", "listed": True, "limit": 5}, 1)),
        _turn(_call("give_answer", {"text": "One midfielder is transfer listed: Marco Midfield, asking £6m. Also £99m.",
                                    "show": [{"result": 1, "title": "Listed midfielders",
                                              "columns": ["player", "club", "asking_price"]}],
                                    "follow_ups": ["Only under 25", "Compare them"]}, 2)),
    ]
    got = await _ask(client, buyer, "show me 5 midfielders who are transfer listed")
    (table,) = got["blocks"]
    assert table["title"] == "Listed midfielders" and table["columns"] == ["player", "club", "asking_price"]
    assert [r["player"] for r in table["rows"]] == ["Marco Midfield"] and table["rows"][0]["path"].startswith("/sales/")
    assert table["source"]["filters"] == {"position": "MID", "listed": True, "sort_by": "value", "sort_dir": "desc"}
    # An invented figure (£99m) is dropped; the real one (£6m) stays.
    assert "£6m" in got["answer"] and "£99m" not in got["answer"]
    assert got["follow_ups"] == ["Only under 25", "Compare them"]
    assert "search_players" in seen[0]["tools"] and "give_answer" in seen[0]["tools"]


async def test_tool_calls_are_capped_and_then_an_answer_is_forced(client: AsyncClient, db, model):
    script, seen = model
    seller, buyer, *_ = await _setup(client, db)
    script += [_turn(_call("money", {}, i)) for i in range(4)]
    script += [_turn(_call("give_answer", {"text": "Here's your budget."}, 9))]
    got = await _ask(client, buyer, "how much do we have?")
    assert len(got["tool_calls"]) == 4
    assert seen[-1]["tool_choice"] == {"type": "function", "function": {"name": "give_answer"}}
    assert got["answer"] == "Here's your budget."


async def test_a_bid_request_opens_the_form_without_the_model(client: AsyncClient, db, model):
    script, seen = model
    from tests.test_lite_buy import _value

    seller, buyer, mid, *_ = await _setup(client, db)
    await _value(db, mid["id"], 6_000_000)
    got = await _ask(client, buyer, "place 6m bid on Marco Midfield")
    assert got["proposal"]["card_path"].startswith(f"/offers/new?player_id={mid['id']}&fee=6000000")
    assert seen == []  # no model call at all


async def test_follow_ups_carry_the_conversation(client: AsyncClient, db, model):
    script, seen = model
    seller, buyer, *_ = await _setup(client, db)
    script += [_turn(_call("give_answer", {"text": "Only one of them is under 25."}, 1))]
    await _ask(client, buyer, "only under 25", history=[{"question": "listed midfielders?", "answer": "Marco Midfield."}])
    roles = [m["role"] for m in seen[0]["messages"]]
    assert roles == ["system", "user", "assistant", "user"]


async def test_a_chart_is_built_from_numeric_columns_only(client: AsyncClient, db, model):
    script, seen = model
    seller, buyer, *_ = await _setup(client, db)
    script += [
        _turn(_call("squad", {}, 1)),
        _turn(_call("give_answer", {"text": "Here they are.",
                                    "charts": [{"result": 1, "x": "player", "y": ["age", "player"], "title": "Ages"},
                                               {"result": 7, "x": "player", "y": ["age"]}]}, 2)),
    ]
    got = await _ask(client, seller, "chart our squad by age")
    charts = [b for b in got["blocks"] if b["type"] == "chart"]
    assert len(charts) == 1 and charts[0]["y"] == ["age"] and charts[0]["title"] == "Ages"  # non-numeric y and a bad result dropped


async def test_context_from_a_page_is_resolved_and_checked(client: AsyncClient, db, model):
    script, seen = model
    seller, buyer, mid, *_ = await _setup(client, db)
    script += [_turn(_call("give_answer", {"text": "OK."}, 1))] * 2
    await client.post("/ai/analyst", json={"question": "tell me about him", "context": {"type": "player", "id": mid["id"]}},
                      headers=_auth_headers(buyer))
    assert "Marco Midfield" in seen[-1]["messages"][0]["content"]
    # Someone else's offer can't be smuggled in as context.
    import uuid as _uuid

    await client.post("/ai/analyst", json={"question": "and this?", "context": {"type": "offer", "id": str(_uuid.uuid4())}},
                      headers=_auth_headers(buyer))
    assert "looking at the offer" not in seen[-1]["messages"][0]["content"]


async def test_a_table_exports_to_excel(client: AsyncClient, db):
    from io import BytesIO

    from openpyxl import load_workbook

    seller, buyer, *_ = await _setup(client, db)
    resp = await client.post("/ai/analyst/export", json={
        "title": "Listed midfielders", "columns": ["player", "asking_price"], "headers": ["Player", "Asking price"],
        "rows": [{"player": "Marco Midfield", "asking_price": 6000000}], "source": "Player market · as of 20:40"},
        headers=_auth_headers(buyer))
    assert resp.status_code == 200 and "listed-midfielders.xlsx" in resp.headers["content-disposition"]
    ws = load_workbook(BytesIO(resp.content)).active
    assert [c.value for c in ws[1]] == ["Player", "Asking price"] and ws["B2"].value == 6000000
